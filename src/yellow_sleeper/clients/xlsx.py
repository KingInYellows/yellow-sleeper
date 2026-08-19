"""Load and write a user value book (csv or xlsx)."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class UserBookRow:
    sleeper_id: str
    name: str
    position: str
    value: float


def load_user_book(path: Path) -> dict[str, UserBookRow]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _load_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _load_xlsx(path)
    else:
        raise ValueError(f"unsupported value book suffix: {suffix or path.name}")
    return {row.sleeper_id: row for row in rows if row.sleeper_id}


def write_user_book(path: Path, rows: list[UserBookRow] | list[dict[str, Any]]) -> None:
    parsed = [_coerce_row(row) for row in rows]
    suffix = path.suffix.lower()
    path.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".csv":
        _write_csv(path, parsed)
        return
    if suffix in {".xlsx", ".xlsm"}:
        _write_xlsx(path, parsed)
        return
    raise ValueError(f"unsupported value book suffix: {suffix or path.name}")


def _coerce_row(row: UserBookRow | dict[str, Any]) -> UserBookRow:
    if isinstance(row, UserBookRow):
        return row
    return UserBookRow(
        sleeper_id=str(row.get("sleeper_id") or ""),
        name=str(row.get("name") or ""),
        position=str(row.get("position") or ""),
        value=float(row["value"]),
    )


def _load_csv(path: Path) -> list[UserBookRow]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return [_row_from_mapping(row) for row in reader if _has_id(row)]


def _write_csv(path: Path, rows: list[UserBookRow]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["sleeper_id", "name", "position", "value"],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "sleeper_id": row.sleeper_id,
                    "name": row.name,
                    "position": row.position,
                    "value": row.value,
                }
            )


def _load_xlsx(path: Path) -> list[UserBookRow]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [str(cell or "").strip().lower() for cell in rows[0]]
    parsed: list[UserBookRow] = []
    for raw in rows[1:]:
        mapping = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        if _has_id(mapping):
            parsed.append(_row_from_mapping(mapping))
    return parsed


def _write_xlsx(path: Path, rows: list[UserBookRow]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "values"
    sheet.append(["sleeper_id", "name", "position", "value"])
    for row in rows:
        sheet.append([row.sleeper_id, row.name, row.position, row.value])
    workbook.save(path)


def _has_id(row: dict[str, Any]) -> bool:
    return bool(str(row.get("sleeper_id") or "").strip())


def _row_from_mapping(row: dict[str, Any]) -> UserBookRow:
    return UserBookRow(
        sleeper_id=str(row.get("sleeper_id") or "").strip(),
        name=str(row.get("name") or "").strip(),
        position=str(row.get("position") or "").strip(),
        value=float(row.get("value") or 0),
    )
